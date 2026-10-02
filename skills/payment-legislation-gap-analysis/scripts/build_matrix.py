#!/usr/bin/env python3
"""Condense per-lane payment-legislation findings JSON into a decision matrix (.xlsx) + executive
summary (.md). Lane-agnostic: each lane defines its own requirement areas.

Usage:
  python build_matrix.py <workspace>            # build xlsx + md
  python build_matrix.py <workspace> --check    # validate findings only
"""
import json, sys, glob, os
from pathlib import Path

RATINGS = {"red", "amber", "green", "unknown"}
RATING_ORDER = {"red": 0, "amber": 1, "unknown": 2, "green": 3}
FILLS = {"red": "F4CCCC", "amber": "FCE8B2", "green": "D9EAD3", "unknown": "E0E0E0", "—": "FFFFFF"}
CTRL = {"present", "partial", "missing", "not_found"}
TEST = {"unit_tested", "integration_only", "untested", "unknown", "fast_tested"}  # fast_tested = legacy spelling, normalised on load
RISK = {"high", "medium", "low"}


def validate(path):
    errs = []
    try:
        doc = json.loads(Path(path).read_text())
    except Exception as e:
        return [f"{Path(path).name}: invalid JSON ({e})"], None
    for a in doc.get("areas", []):
        for f in a.get("findings", []):
            if f.get("test_status") == "fast_tested":
                f["test_status"] = "unit_tested"  # legacy spelling -> canonical
    for key in ("lane", "lane_name", "areas"):
        if key not in doc:
            errs.append(f"{Path(path).name}: missing '{key}'")
    seen = set()
    for a in doc.get("areas", []):
        area = a.get("area")
        if not area:
            errs.append(f"{Path(path).name}: area with no name")
        if area in seen:
            errs.append(f"{Path(path).name}: duplicate area '{area}'")
        seen.add(area)
        if a.get("rating") not in RATINGS:
            errs.append(f"{Path(path).name}: area '{area}' bad rating '{a.get('rating')}'")
        for f in a.get("findings", []):
            if f.get("control_status") not in CTRL:
                errs.append(f"{Path(path).name}: {f.get('id')} bad control_status '{f.get('control_status')}'")
            if f.get("test_status") not in TEST:
                errs.append(f"{Path(path).name}: {f.get('id')} bad test_status '{f.get('test_status')}'")
            if f.get("risk") not in RISK:
                errs.append(f"{Path(path).name}: {f.get('id')} bad risk '{f.get('risk')}'")
    return errs, doc


def overall(doc):
    ratings = [a.get("rating", "unknown") for a in doc.get("areas", [])]
    return min(ratings, key=lambda r: RATING_ORDER.get(r, 2)) if ratings else "unknown"


def counts(doc):
    c = {"high": 0, "medium": 0, "low": 0}
    for a in doc.get("areas", []):
        for f in a.get("findings", []):
            r = f.get("risk")
            if r in c:
                c[r] += 1
    return c


def load(workspace):
    docs, errs = [], []
    for p in sorted(glob.glob(os.path.join(workspace, "findings", "*.json"))):
        e, d = validate(p)
        errs += e
        if d:
            docs.append(d)
    return docs, errs


def build_xlsx(docs, out):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    wb = Workbook()

    # Sheet 1: lane summary
    ws = wb.active
    ws.title = "Lanes"
    ws.append(["Lane", "Name", "Overall", "High", "Medium", "Low", "Areas (area:rating)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for doc in sorted(docs, key=lambda d: (RATING_ORDER.get(overall(d), 2), d.get("lane", ""))):
        cc = counts(doc)
        ov = overall(doc)
        areas = "; ".join(f"{a['area']}:{a['rating']}" for a in doc.get("areas", []))
        ws.append([doc.get("lane"), doc.get("lane_name"), ov, cc["high"], cc["medium"], cc["low"], areas])
        ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=FILLS.get(ov, "FFFFFF"))
    for col, w in {"A": 14, "B": 34, "C": 9, "D": 6, "E": 8, "F": 6, "G": 80}.items():
        ws.column_dimensions[col].width = w

    # Sheet 2: all findings (filterable)
    wf = wb.create_sheet("Findings")
    hdr = ["Risk", "Lane", "Area", "ID", "Requirement", "Citation", "Conf", "Control", "Test", "Evidence", "Recommendation"]
    wf.append(hdr)
    for c in wf[1]:
        c.font = Font(bold=True)
    rows = []
    for doc in docs:
        for a in doc.get("areas", []):
            for f in a.get("findings", []):
                ev = " | ".join(f"{e.get('file','')}:{e.get('lines','')} {e.get('note','')}" for e in f.get("evidence", []))
                rows.append([f.get("risk"), doc.get("lane"), a.get("area"), f.get("id"), f.get("requirement"),
                             f.get("citation"), f.get("citation_confidence"), f.get("control_status"),
                             f.get("test_status"), ev, f.get("recommendation")])
    rows.sort(key=lambda r: ({"high": 0, "medium": 1, "low": 2}.get(r[0], 3), r[1] or "", r[3] or ""))
    for r in rows:
        wf.append(r)
    wf.auto_filter.ref = f"A1:K{wf.max_row}"
    for col, w in {"A": 8, "B": 12, "C": 22, "D": 14, "E": 46, "F": 26, "G": 6, "H": 10, "I": 16, "J": 50, "K": 50}.items():
        wf.column_dimensions[col].width = w
    for row in wf.iter_rows(min_row=2, max_col=1):
        row[0].fill = PatternFill("solid", fgColor={"high": "F4CCCC", "medium": "FCE8B2", "low": "D9EAD3"}.get(row[0].value, "FFFFFF"))

    wb.save(out)


def build_summary(docs, out):
    total = {"high": 0, "medium": 0, "low": 0}
    for doc in docs:
        c = counts(doc)
        for k in total:
            total[k] += c[k]
    red = [d.get("lane") for d in docs if overall(d) == "red"]
    # code-gap vs test-gap split
    missing = present_untested = green_over_wrong = 0
    for doc in docs:
        for a in doc.get("areas", []):
            for f in a.get("findings", []):
                if f.get("control_status") in ("missing", "not_found"):
                    missing += 1
                elif f.get("control_status") in ("present", "partial") and f.get("test_status") in ("untested", "unknown"):
                    present_untested += 1
    lines = []
    lines.append("# Payment Legislation Gap Analysis — Executive Summary\n")
    lines.append(f"*Generated from {len(docs)} legislation lanes.*\n")
    lines.append("> ORCHESTRATOR: replace this note with 2-3 paragraphs of synthesis — the cross-lane "
                 "patterns, the code-gap vs test-gap split, and what to fix first. Stats below are computed.\n")
    lines.append("## Headline numbers\n")
    lines.append(f"- Findings: **{sum(total.values())}** total — **{total['high']} high**, {total['medium']} medium, {total['low']} low.")
    lines.append(f"- Lanes rated red overall: **{', '.join(sorted(red)) or 'none'}**.")
    lines.append(f"- Code gaps (missing/not_found controls): **{missing}**. Present-but-untested controls: **{present_untested}**.\n")
    lines.append("## Top high-risk findings\n")
    hi = []
    for doc in docs:
        for a in doc.get("areas", []):
            for f in a.get("findings", []):
                if f.get("risk") == "high":
                    hi.append((doc.get("lane"), a.get("area"), f))
    for lane, area, f in hi[:12]:
        lines.append(f"- **{f.get('id')}** ({lane}, {area}): {f.get('requirement')} — control `{f.get('control_status')}`, test `{f.get('test_status')}`. {f.get('recommendation','')}")
    if len(hi) > 12:
        lines.append(f"- …and {len(hi)-12} more high-risk findings (see Findings sheet).\n")
    lines.append("\n## Caveats\n")
    for doc in docs:
        cav = doc.get("confidence_caveats")
        if cav:
            lines.append(f"- **{doc.get('lane')}**: {cav}")
        refreshed = doc.get("legislation_refreshed", "")
        if refreshed and refreshed != "not-refreshed":
            lines.append(f"  - *legislation:* {refreshed}")
    lines.append("\n*Not legal advice. Citations researched/refreshed as of the run date and spot-verified; confirm with compliance/legal before external use.*")
    Path(out).write_text("\n".join(lines))


def main():
    if len(sys.argv) < 2:
        print("usage: build_matrix.py <workspace> [--check]"); sys.exit(2)
    workspace = sys.argv[1]
    check = "--check" in sys.argv[2:]
    docs, errs = load(workspace)
    if errs:
        print("VALIDATION ERRORS:")
        for e in errs:
            print("  " + e)
    else:
        print(f"All {len(docs)} findings files valid.")
    if check:
        sys.exit(1 if errs else 0)
    if not docs:
        print("No valid findings to build."); sys.exit(1)
    build_xlsx(docs, os.path.join(workspace, "decision_matrix.xlsx"))
    build_summary(docs, os.path.join(workspace, "executive_summary.md"))
    print(f"Wrote {os.path.join(workspace,'decision_matrix.xlsx')} and {os.path.join(workspace,'executive_summary.md')}")


if __name__ == "__main__":
    main()
